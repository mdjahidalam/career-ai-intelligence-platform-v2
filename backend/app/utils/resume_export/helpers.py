"""
==================================================
Resume Helper Utilities
Career AI Intelligence Platform
==================================================
"""

from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .constants import *


class ResumeHelper:

    # ==================================================
    # Apply Page Layout
    # ==================================================

    @staticmethod
    def apply_page_margin(document):

        section = document.sections[0]

        section.top_margin = Inches(PAGE_MARGIN)
        section.bottom_margin = Inches(PAGE_MARGIN)

        section.left_margin = Inches(PAGE_MARGIN)
        section.right_margin = Inches(PAGE_MARGIN)

    # ==================================================
    # Document Properties
    # ==================================================

    @staticmethod
    def set_document_properties(document):

        properties = document.core_properties

        properties.author = "Career AI Intelligence Platform"

        properties.company = "Career AI Intelligence Platform"

        properties.title = "Professional Resume"

        properties.subject = "ATS Friendly Resume"

        properties.comments = (
            "Automatically generated using Career AI Intelligence Platform"
        )

        properties.category = "Resume"

    # ==================================================
    # Professional Divider
    # ==================================================

    @staticmethod
    def divider(document):

        paragraph = document.add_paragraph()

        pPr = paragraph._p.get_or_add_pPr()

        border = OxmlElement("w:pBdr")

        bottom = OxmlElement("w:bottom")

        bottom.set(qn("w:val"), "single")

        bottom.set(qn("w:sz"), "8")

        bottom.set(qn("w:space"), "1")

        bottom.set(qn("w:color"), DIVIDER_COLOR)

        border.append(bottom)

        pPr.append(border)

    # ==================================================
    # Empty Space
    # ==================================================

    @staticmethod
    def spacer(

        document,

        count=1

    ):

        for _ in range(count):

            document.add_paragraph()

    # ==================================================
    # Safe Value
    # ==================================================

    @staticmethod
    def safe(value):

        if value is None:

            return ""

        return str(value).strip()

    # ==================================================
    # Join Values
    # ==================================================

    @staticmethod
    def join(

        values,

        separator=" | "

    ):

        cleaned = [

            ResumeHelper.safe(item)

            for item in values

            if ResumeHelper.safe(item)

        ]

        return separator.join(cleaned)

    # ==================================================
    # Bullet List
    # ==================================================

    @staticmethod
    def add_bullets(

        document,

        items

    ):

        if not items:

            return

        for item in items:

            paragraph = document.add_paragraph(

                style="List Bullet"

            )

            paragraph.add_run(

                ResumeHelper.safe(item)

            )

    # ==================================================
    # Key : Value
    # ==================================================

    @staticmethod
    def key_value(

        document,

        key,

        value

    ):

        if not value:

            return

        paragraph = document.add_paragraph()

        paragraph.add_run(

            key

        ).bold = True

        paragraph.add_run(

            ResumeHelper.safe(value)

        )

    # ==================================================
    # Horizontal Space
    # ==================================================

    @staticmethod
    def horizontal_space(

        paragraph,

        count=3

    ):

        paragraph.add_run(

            " " * count

        )
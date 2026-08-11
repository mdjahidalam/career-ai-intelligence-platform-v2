"""
==================================================
Resume Sections
Career AI Intelligence Platform
Microsoft-Style Resume DOCX Renderer
==================================================
"""

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

from .styles import ResumeStyle
from .hyperlinks import ResumeHyperlink
from .helpers import ResumeHelper


class ResumeSections:

    # ==================================================
    # Header
    # ==================================================

    @staticmethod
    def header(
        document,
        candidate: dict
    ):

        candidate = candidate or {}

        # ------------------------------------------
        # Name + Role
        # ------------------------------------------

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_PARAGRAPH_ALIGNMENT.CENTER
        )

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0

        name = str(
            candidate.get(
                "name",
                "Candidate Name"
            ) or "Candidate Name"
        ).strip()

        role = str(
            candidate.get(
                "role",
                ""
            ) or ""
        ).strip()

        name_run = paragraph.add_run(name)

        ResumeStyle.name(name_run)

        if role:

            role_run = paragraph.add_run(
                f" ({role})"
            )

            ResumeStyle.title(role_run)

        # ------------------------------------------
        # Contact Line
        # ------------------------------------------

        contact_items = []

        email = str(
            candidate.get("email", "") or ""
        ).strip()

        phone = str(
            candidate.get("phone", "") or ""
        ).strip()

        location = str(
            candidate.get("location", "") or ""
        ).strip()

        if email:
            contact_items.append(
                ("email", email)
            )

        if phone:
            contact_items.append(
                ("phone", phone)
            )

        if location:
            contact_items.append(
                ("text", location)
            )

        if contact_items:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_PARAGRAPH_ALIGNMENT.CENTER
            )

            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0

            for index, (item_type, value) in enumerate(
                contact_items
            ):

                if index > 0:

                    separator = paragraph.add_run(
                        " | "
                    )

                    ResumeStyle.small(
                        separator
                    )

                if item_type == "email":

                    ResumeHyperlink.email(
                        paragraph,
                        value
                    )

                elif item_type == "phone":

                    ResumeHyperlink.phone(
                        paragraph,
                        value
                    )

                else:

                    run = paragraph.add_run(
                        value
                    )

                    ResumeStyle.small(
                        run
                    )

        # ------------------------------------------
        # Social / Professional Links
        # ------------------------------------------

        links = []

        linkedin = str(
            candidate.get("linkedin", "") or ""
        ).strip()

        github = str(
            candidate.get("github", "") or ""
        ).strip()

        portfolio = str(
            candidate.get("portfolio", "") or ""
        ).strip()

        if linkedin:
            links.append(
                ("LinkedIn", linkedin)
            )

        if github:
            links.append(
                ("GitHub", github)
            )

        if portfolio:
            links.append(
                ("Portfolio", portfolio)
            )

        if links:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_PARAGRAPH_ALIGNMENT.CENTER
            )

            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.0

            for index, (label, url) in enumerate(
                links
            ):

                if index > 0:

                    separator = paragraph.add_run(
                        " | "
                    )

                    ResumeStyle.small(
                        separator
                    )

                ResumeHyperlink.add(
                    paragraph,
                    label,
                    url
                )

        # ------------------------------------------
        # Divider
        # ------------------------------------------

        ResumeHelper.divider(
            document
        )

    # ==================================================
    # Section Title
    # ==================================================

    @staticmethod
    def section_title(
        document,
        title
    ):

        if not title:
            return

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(
            str(title).upper()
        )

        ResumeStyle.section(
            run
        )

    # ==================================================
    # Professional Summary
    # ==================================================

    @staticmethod
    def summary(
        document,
        summary
    ):

        if not summary:
            return

        ResumeSections.section_title(
            document,
            "Professional Summary"
        )

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(
            str(summary).strip()
        )

        ResumeStyle.body(
            run
        )

    # ==================================================
    # Technical Skills
    # ==================================================

    @staticmethod
    def skills(
        document,
        skills: dict
    ):

        if not skills:
            return

        ResumeSections.section_title(
            document,
            "Technical Skills"
        )

        for category, values in skills.items():

            if not values:
                continue

            # --------------------------------------
            # Normalize values
            # --------------------------------------

            if not isinstance(
                values,
                list
            ):

                values = [values]

            cleaned = []

            for value in values:

                if value is None:
                    continue

                value = str(value).strip()

                if not value:
                    continue

                if value not in cleaned:
                    cleaned.append(value)

            if not cleaned:
                continue

            # --------------------------------------
            # Compact inline format
            # --------------------------------------

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0

            category_run = paragraph.add_run(
                f"{str(category).strip()}: "
            )

            category_run.bold = True

            ResumeStyle.body(
                category_run
            )

            skills_run = paragraph.add_run(
                ", ".join(cleaned)
            )

            ResumeStyle.body(
                skills_run
            )

    # ==================================================
    # Experience
    # ==================================================

    @staticmethod
    def experience(
        document,
        experiences: list
    ):

        if not experiences:
            return

        ResumeSections.section_title(
            document,
            "Experience"
        )

        for exp in experiences:

            if not isinstance(exp, dict):
                continue

            role = str(
                exp.get("role", "") or ""
            ).strip()

            company = str(
                exp.get("company", "") or ""
            ).strip()

            location = str(
                exp.get("location", "") or ""
            ).strip()

            duration = str(
                exp.get("duration", "") or ""
            ).strip()

            # --------------------------------------
            # Role
            # --------------------------------------

            if role:

                paragraph = document.add_paragraph()

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0

                run = paragraph.add_run(
                    role
                )

                # run.bold = True

                ResumeStyle.body(
                    run
                )

            # --------------------------------------
            # Company / Location / Duration
            # --------------------------------------

            details = [
                value
                for value in [
                    company,
                    location,
                    duration
                ]
                if value
            ]

            if details:

                paragraph = document.add_paragraph()

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0

                run = paragraph.add_run(
                    " | ".join(details)
                )

                ResumeStyle.small(
                    run
                )

            # --------------------------------------
            # Bullet Points
            # --------------------------------------

            points = exp.get(
                "highlights",
                exp.get(
                    "bullets",
                    exp.get(
                        "description",
                        []
                    )
                )
            )

            ResumeSections._add_bullets(
                document,
                points
            )

    # ==================================================
    # Projects
    # ==================================================

    @staticmethod
    def projects(
        document,
        projects: list
    ):

        if not projects:
            return

        ResumeSections.section_title(
            document,
            "Projects"
        )

        for project in projects:

            if not isinstance(project, dict):
                continue

            # --------------------------------------
            # Project Title
            # --------------------------------------

            title = str(
                project.get(
                    "title",
                    project.get(
                        "name",
                        ""
                    )
                ) or ""
            ).strip()

            if title:

                paragraph = document.add_paragraph()

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0

                run = paragraph.add_run(
                    title
                )

                # run.bold = True

                ResumeStyle.body(
                    run
                )

            # --------------------------------------
            # Project Bullets
            # --------------------------------------

            points = project.get(
                "highlights",
                project.get(
                    "bullets",
                    project.get(
                        "description",
                        []
                    )
                )
            )

            ResumeSections._add_bullets(
                document,
                points
            )

            # --------------------------------------
            # Technologies
            # --------------------------------------

            technologies = project.get(
                "technologies",
                []
            )

            if technologies:

                if not isinstance(
                    technologies,
                    list
                ):

                    technologies = [
                        technologies
                    ]

                technologies = [
                    str(item).strip()
                    for item in technologies
                    if str(item).strip()
                ]

                if technologies:

                    paragraph = document.add_paragraph()

                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.0

                    label = paragraph.add_run(
                        "Technologies: "
                    )

                    label.bold = True

                    ResumeStyle.small(
                        label
                    )

                    run = paragraph.add_run(
                        ", ".join(technologies)
                    )

                    ResumeStyle.small(
                        run
                    )

            # --------------------------------------
            # Impact
            # --------------------------------------

            impact = str(
                project.get(
                    "impact",
                    ""
                ) or ""
            ).strip()

            if impact:

                paragraph = document.add_paragraph()

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0

                label = paragraph.add_run(
                    "Impact: "
                )

                label.bold = True

                ResumeStyle.small(
                    label
                )

                run = paragraph.add_run(
                    impact
                )

                ResumeStyle.small(
                    run
                )

    # ==================================================
    # Education
    # ==================================================

    @staticmethod
    def education(
        document,
        education: list
    ):

        if not education:
            return

        ResumeSections.section_title(
            document,
            "Education"
        )

        for item in education:

            if not isinstance(item, dict):
                continue

            # --------------------------------------
            # Degree
            # --------------------------------------

            degree = str(
                item.get(
                    "degree",
                    ""
                ) or ""
            ).strip()

            if degree:

                paragraph = document.add_paragraph()

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0

                run = paragraph.add_run(
                    degree
                )

                # run.bold = True

                ResumeStyle.body(
                    run
                )

            # --------------------------------------
            # Education Details
            # --------------------------------------

            details = []

            specialization = str(
                item.get(
                    "specialization",
                    ""
                ) or ""
            ).strip()

            institute = str(
                item.get(
                    "institute",
                    ""
                ) or ""
            ).strip()

            year = str(
                item.get(
                    "year",
                    ""
                ) or ""
            ).strip()

            score = str(
                item.get(
                    "cgpa",
                    ""
                ) or ""
            ).strip()

            score_type = str(
                item.get(
                    "score_type",
                    ""
                ) or ""
            ).strip()

            if specialization:
                details.append(
                    specialization
                )

            if institute:
                details.append(
                    institute
                )

            if score:

                if score_type:

                    details.append(
                        f"{score_type}: {score}"
                    )

                else:

                    details.append(
                        f"CGPA: {score}"
                    )

            if year:
                details.append(
                    year
                )

            if details:

                paragraph = document.add_paragraph()

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0

                run = paragraph.add_run(
                    " | ".join(details)
                )

                ResumeStyle.small(
                    run
                )

    # ==================================================
    # Certifications
    # ==================================================

    @staticmethod
    def certifications(
        document,
        certifications: list
    ):

        if not certifications:
            return

        ResumeSections.section_title(
            document,
            "Certifications"
        )

        for cert in certifications:

            if isinstance(
                cert,
                dict
            ):

                name = str(
                    cert.get(
                        "name",
                        ""
                    ) or ""
                ).strip()

                issuer = str(
                    cert.get(
                        "issuer",
                        ""
                    ) or ""
                ).strip()

                year = str(
                    cert.get(
                        "year",
                        ""
                    ) or ""
                ).strip()

                parts = []

                if name:
                    parts.append(name)

                if issuer:
                    parts.append(issuer)

                if year:
                    parts.append(year)

                text = " | ".join(parts)

            else:

                text = str(
                    cert
                ).strip()

            if not text:
                continue

            ResumeSections._add_single_bullet(
                document,
                text
            )

    # ==================================================
    # Achievements
    # ==================================================

    @staticmethod
    def achievements(
        document,
        achievements: list
    ):

        if not achievements:
            return

        ResumeSections.section_title(
            document,
            "Achievements"
        )

        for achievement in achievements:

            if isinstance(
                achievement,
                dict
            ):

                title = str(
                    achievement.get(
                        "title",
                        ""
                    ) or ""
                ).strip()

                description = str(
                    achievement.get(
                        "description",
                        ""
                    ) or ""
                ).strip()

                if title and description:

                    text = (
                        f"{title} - "
                        f"{description}"
                    )

                else:

                    text = (
                        title or
                        description
                    )

            else:

                text = str(
                    achievement
                ).strip()

            if not text:
                continue

            ResumeSections._add_single_bullet(
                document,
                text
            )

    # ==================================================
    # Languages
    # ==================================================

    @staticmethod
    def languages(
        document,
        languages: list
    ):

        if not languages:
            return

        ResumeSections.section_title(
            document,
            "Languages"
        )

        values = []

        for language in languages:

            if isinstance(
                language,
                dict
            ):

                name = str(
                    language.get(
                        "name",
                        ""
                    ) or ""
                ).strip()

                level = str(
                    language.get(
                        "level",
                        ""
                    ) or ""
                ).strip()

                if name and level:

                    values.append(
                        f"{name} ({level})"
                    )

                elif name:

                    values.append(name)

            else:

                value = str(
                    language
                ).strip()

                if value:
                    values.append(value)

        if not values:
            return

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(
            " | ".join(values)
        )

        ResumeStyle.body(
            run
        )

    # ==================================================
    # Bullet Helpers
    # ==================================================

    @staticmethod
    def _add_bullets(
        document,
        points
    ):

        if not points:
            return

        if not isinstance(
            points,
            list
        ):

            points = [points]

        for point in points:

            if isinstance(
                point,
                dict
            ):

                point = (
                    point.get("text")
                    or point.get("description")
                    or point.get("content")
                    or ""
                )

            point = str(
                point
            ).strip()

            if not point:
                continue

            ResumeSections._add_single_bullet(
                document,
                point
            )

    @staticmethod
    def _add_single_bullet(
        document,
        text
    ):

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.paragraph_format.left_indent = Pt(12)
        paragraph.paragraph_format.first_line_indent = Pt(-6)

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(
            str(text).strip()
        )

        ResumeStyle.body(
            run
        )
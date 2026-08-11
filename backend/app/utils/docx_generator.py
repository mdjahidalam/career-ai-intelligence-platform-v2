from io import BytesIO

from docx import Document

from docx.shared import Pt


class ResumeDOCXGenerator:

    @staticmethod
    def generate(data: dict):

        document = Document()

        # ==========================================
        # Candidate
        # ==========================================

        candidate = data.get("candidate", {})

        name = candidate.get("name", "Candidate")

        document.add_heading(name, level=1)

        details = []

        if candidate.get("email"):
            details.append(candidate["email"])

        if candidate.get("phone"):
            details.append(candidate["phone"])

        if candidate.get("linkedin"):
            details.append(candidate["linkedin"])

        if candidate.get("github"):
            details.append(candidate["github"])

        if candidate.get("portfolio"):
            details.append(candidate["portfolio"])

        p = document.add_paragraph()

        p.style.font.size = Pt(11)

        p.add_run(" | ".join(details))

        # ==========================================
        # Summary
        # ==========================================

        document.add_heading(
            "Professional Summary",
            level=2
        )

        document.add_paragraph(

            data.get(
                "professional_summary",
                ""
            )

        )

        # ==========================================
        # Skills
        # ==========================================

        document.add_heading(

            "Technical Skills",

            level=2

        )

        skills = data.get(

            "skills",

            {}

        )

        for category, values in skills.items():

            if values:

                document.add_paragraph(

                    f"{category.title()} : "

                    + ", ".join(values)

                )

        # ==========================================
        # Experience
        # ==========================================

        experience = data.get(

            "experience",

            []

        )

        if experience:

            document.add_heading(

                "Experience",

                level=2

            )

            for exp in experience:

                document.add_heading(

                    exp.get(

                        "role",

                        ""

                    ),

                    level=3

                )

                document.add_paragraph(

                    exp.get(

                        "company",

                        ""

                    )

                )

                document.add_paragraph(

                    exp.get(

                        "description",

                        ""

                    )

                )

        # ==========================================
        # Projects
        # ==========================================

        document.add_heading(

            "Projects",

            level=2

        )

        for project in data.get(

            "projects",

            []

        ):

            document.add_heading(

                project.get(

                    "title",

                    ""

                ),

                level=3

            )

            document.add_paragraph(

                project.get(

                    "description",

                    ""

                )

            )

            technologies = project.get(

                "technologies",

                []

            )

            if technologies:

                document.add_paragraph(

                    "Technologies : "

                    +

                    ", ".join(

                        technologies

                    )

                )

        # ==========================================
        # Education
        # ==========================================

        document.add_heading(

            "Education",

            level=2

        )

        for edu in data.get(

            "education",

            []

        ):

            document.add_paragraph(

                f"{edu.get('degree')} "

                f"{edu.get('branch')}"

            )

            document.add_paragraph(

                edu.get(

                    "college",

                    ""

                )

            )

            document.add_paragraph(

                f"{edu.get('cgpa')}"

                f" | "

                f"{edu.get('year')}"

            )

        # ==========================================
        # Certifications
        # ==========================================

        certs = data.get(

            "certifications",

            []

        )

        if certs:

            document.add_heading(

                "Certifications",

                level=2

            )

            for cert in certs:

                document.add_paragraph(

                    cert,

                    style="List Bullet"

                )

        # ==========================================
        # Save
        # ==========================================

        file_stream = BytesIO()

        document.save(

            file_stream

        )

        file_stream.seek(0)

        return file_stream